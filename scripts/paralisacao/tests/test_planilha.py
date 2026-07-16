
import openpyxl
import pytest

from paralisacao.planilha import ler_planilha


@pytest.fixture
def planilha_teste(tmp_path):
    path = tmp_path / "planilha_teste.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active

    # Linhas de metadados
    ws.append(["Programa do Navio", "1766"])
    ws.append([])

    # Cabeçalho
    ws.append(["Registro", "Navio", "Data", "Produto", "Período", "Início", "Fim", "Motivo"])

    # Dados de teste
    ws.append(["1", "YANGTZE QUANTUM", "09/07/2026", "Milho", "13:00/19:00", "13:00", "13:15", "DDS"])
    ws.append(["1", "YANGTZE QUANTUM", "09/07/2026", "Milho", "13:00/19:00", "16:00", "16:25", "Defeito no Speed"])

    wb.save(path)
    return path

class TestLerPlanilha:
    def test_leitura_completa(self, planilha_teste):
        dados = ler_planilha(str(planilha_teste))

        assert dados.prog_navio == "1766"
        assert len(dados.capas) == 1

        capa = dados.capas[0]
        assert capa.registro == "1"
        assert capa.navio == "YANGTZE QUANTUM"
        assert capa.data == "09/07/2026"
        assert capa.produto == "Milho"
        assert capa.periodo == "13:00/19:00"

        assert len(capa.paralisacoes) == 2
        assert capa.paralisacoes[0].inicio == "13:00"
        assert capa.paralisacoes[0].fim == "13:15"
        assert capa.paralisacoes[0].motivo == "DDS"

        assert capa.paralisacoes[1].inicio == "16:00"
        assert capa.paralisacoes[1].fim == "16:25"
        assert capa.paralisacoes[1].motivo == "Defeito no Speed"

    def test_planilha_invalida_sem_cabecalho(self, tmp_path):
        path = tmp_path / "planilha_invalida.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Coluna1", "Coluna2"])
        wb.save(path)

        with pytest.raises(ValueError, match="Cabeçalho da tabela.*não encontrado"):
            ler_planilha(str(path))

    def test_planilha_invalida_coluna_faltando(self, tmp_path):
        path = tmp_path / "planilha_colunas_faltantes.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Registro", "Navio", "Data", "Motivo"])  # faltam várias colunas
        wb.save(path)

        with pytest.raises(ValueError, match="Colunas ausentes na planilha"):
            ler_planilha(str(path))
