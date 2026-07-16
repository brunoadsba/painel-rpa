from dataclasses import dataclass, field

import openpyxl


@dataclass
class Paralisacao:
    inicio: str
    fim: str
    motivo: str


@dataclass
class Capa:
    registro: str
    navio: str
    data: str
    produto: str
    periodo: str
    paralisacoes: list[Paralisacao] = field(default_factory=list)


@dataclass
class DadosPlanilha:
    prog_navio: str | None
    capas: list[Capa]


COLUNAS_ESPERADAS = (
    "Registro",
    "Navio",
    "Data",
    "Produto",
    "Período",
    "Início",
    "Fim",
    "Motivo",
)


def _texto(valor) -> str:
    if valor is None:
        return ""
    return str(valor).strip()


def _achar_prog_navio(ws) -> str | None:
    for row in ws.iter_rows(values_only=True):
        rotulo = _texto(row[0]).lower()
        if "programa" in rotulo and "navio" in rotulo:
            return _texto(row[1]) or None
    return None


def _achar_linha_cabecalho(ws) -> int | None:
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        valores = {_texto(v) for v in row}
        if "Registro" in valores and "Motivo" in valores:
            return idx
    return None


def ler_planilha(caminho: str) -> DadosPlanilha:
    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = wb.active

    prog_navio = _achar_prog_navio(ws)
    linha_cabecalho = _achar_linha_cabecalho(ws)
    if linha_cabecalho is None:
        raise ValueError(
            "Cabeçalho da tabela (Registro..Motivo) não encontrado na planilha."
        )

    header_row = list(ws.iter_rows(min_row=linha_cabecalho, max_row=linha_cabecalho, values_only=True))[0]
    col = {_texto(nome): i for i, nome in enumerate(header_row)}
    faltando = [c for c in COLUNAS_ESPERADAS if c not in col]
    if faltando:
        raise ValueError(f"Colunas ausentes na planilha: {', '.join(faltando)}")

    capas: dict[str, Capa] = {}
    ordem: list[str] = []

    for row in ws.iter_rows(min_row=linha_cabecalho + 1, values_only=True):
        registro = _texto(row[col["Registro"]])
        motivo = _texto(row[col["Motivo"]])
        if not registro and not motivo:
            continue

        if registro not in capas:
            capas[registro] = Capa(
                registro=registro,
                navio=_texto(row[col["Navio"]]),
                data=_texto(row[col["Data"]]),
                produto=_texto(row[col["Produto"]]),
                periodo=_texto(row[col["Período"]]),
            )
            ordem.append(registro)

        capas[registro].paralisacoes.append(
            Paralisacao(
                inicio=_texto(row[col["Início"]]),
                fim=_texto(row[col["Fim"]]),
                motivo=motivo,
            )
        )

    return DadosPlanilha(prog_navio=prog_navio, capas=[capas[r] for r in ordem])


if __name__ == "__main__":
    import sys

    from paralisacao import config

    caminho = sys.argv[1] if len(sys.argv) > 1 else config.PLANILHA
    dados = ler_planilha(caminho)
    print(f"Programação do navio: {dados.prog_navio}")
    print(f"Total de capas: {len(dados.capas)}\n")
    for capa in dados.capas:
        print(f"Capa {capa.registro} | {capa.navio} | {capa.data} | período {capa.periodo}")
        for p in capa.paralisacoes:
            print(f"   - {p.inicio} -> {p.fim} | {p.motivo}")
